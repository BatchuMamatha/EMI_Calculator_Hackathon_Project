"""Regenerate docs/Project_Documentation.xlsx in the Finding_Hospital reference style.

Sheets produced (in order):
  UserStories    — title row + blue header row + 10 user stories
  TestScenarios  — plain header + 10 scenarios
  TestCases      — plain header + 10 test cases (multi-row, with merged cells
                   on Group / TC ID / Description / Designer / Type / Data /
                   Complexity across step rows)
  RTM            — plain header + traceability rows
  Defects        — plain header + defects log
"""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = r"C:\Users\2485860\Downloads\EMI_Calculator_Hackathon_Project\docs\Project_Documentation.xlsx"

# Reference style colours
BLUE_HEADER_FILL = "FF2F78C4"
WHITE = "FFFFFFFF"
BLACK = "FF000000"
GREY_BORDER = "FFBFBFBF"

thin = Side(border_style="thin", color=GREY_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def title_font(size=14):
    return Font(name="Arial", size=size, bold=True, color=WHITE)

def header_font():
    return Font(name="Arial", size=9, bold=True, color=WHITE)

def plain_header_font():
    return Font(name="Arial", size=10, bold=True, color=BLACK)

def data_font():
    return Font(name="Aptos Narrow", size=11, color=BLACK)


def style_blue_header_row(ws, row, ncols):
    """Apply the User-Stories blue header style to a row."""
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = PatternFill("solid", fgColor=BLUE_HEADER_FILL)
        cell.font = header_font()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 38


def style_plain_header_row(ws, row, ncols):
    """Apply the Test-Cases plain bold header style."""
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = plain_header_font()
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 24


def style_data_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        if cell.font is None or cell.font.name != "Aptos Narrow":
            cell.font = data_font()
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        cell.border = BORDER


def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


wb = Workbook()
wb.remove(wb.active)

# =====================================================================
# 1) UserStories  (blue header style)
# =====================================================================
us_headers = [
    "User Story ID", "User Story Name",
    "Feature ID\n<Feature for which user stories are created>",
    "As a/an", "I want to", "so that",
    "Author", "Created on", "Reviewed by", "Priority",
    "Predecessor User Story", "Successor User Story", "Acceptance Criteria",
]

us_rows = [
    ["US_001", "Calculate Car Loan EMI", "FI_001",
     "Car loan applicant",
     "Calculate the EMI for a car loan of Rs.15 Lakh at 9.5% interest for 1 year tenure",
     "I know the monthly payment commitment before signing the loan agreement",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "High",
     "", "US_002",
     "1. Open emicalculator.net\n2. Switch to Car Loan tab\n3. Enter 1500000 / 9.5 / 1 year\n4. EMI should be Rs.1,31,524 (±Rs.2 tolerance)"],

    ["US_002", "View First Month Interest and Principal Split", "FI_001",
     "Car loan applicant",
     "View the principal vs interest break-up of the first EMI",
     "I understand how much of my first payment is interest versus principal",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "High",
     "US_001", "US_003",
     "1. After EMI is calculated, expand the year row in the schedule\n2. Read the first month row\n3. Interest = Rs.11,875, Principal = Rs.1,19,649 (±Rs.2)"],

    ["US_003", "Export Car Loan EMI Summary to Excel", "FI_001",
     "Finance team analyst",
     "Export the calculated car loan EMI summary into an Excel workbook",
     "I can share the numbers with stakeholders for offline review",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "Medium",
     "US_002", "US_004",
     "1. Capture EMI, first month split, total interest\n2. Write to output/CarLoan_EMI_Summary.xlsx via Apache POI\n3. Excel file exists with at least 2 rows"],

    ["US_004", "Navigate to Home Loan Calculator via Menu", "FI_002",
     "Home loan applicant",
     "Open the dedicated Home Loan EMI Calculator page from the top menu",
     "I can access the page tailored to home loans",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "Medium",
     "US_003", "US_005",
     "1. Hover 'Loan Calculators & Widgets'\n2. Click 'Home Loan EMI Calculator'\n3. URL contains 'home-loan-emi-calculator'"],

    ["US_005", "Extract Year-on-Year Payment Schedule", "FI_002",
     "Home loan applicant",
     "Fill the home loan form and extract the year-on-year amortisation table",
     "I can plan tax benefits and prepayments year by year",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "High",
     "US_004", "US_006",
     "1. Fill amount 25L / rate 8.5 / tenure 20 yr\n2. Schedule renders\n3. At least 10 yearly rows visible with all 6 columns"],

    ["US_006", "Persist Yearly Schedule to Excel", "FI_002",
     "Home loan applicant",
     "Save the extracted year-on-year schedule into an Excel file",
     "I can compare the amortisation against quotes from multiple banks",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "Medium",
     "US_005", "US_007",
     "1. Write the schedule grid via Apache POI\n2. output/HomeLoan_YearlySchedule.xlsx exists\n3. Sheet 'YearlySchedule' has header + all yearly rows"],

    ["US_007", "Validate Loan Calculator UI Inputs", "FI_003",
     "QA tester",
     "Verify text boxes and sliders on the EMI Calculator sub-tab are enabled and rendered",
     "Users can edit every field and the controls are operable",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "High",
     "US_006", "US_008",
     "1. Open Loan Calculator > EMI Calculator sub-tab\n2. Amount, Interest, Tenure text boxes are enabled\n3. All three sliders rendered"],

    ["US_008", "Tenure Unit Toggle Changes Slider Scale", "FI_003",
     "Loan calculator user",
     "Switch tenure between Year and Month and see the slider scale update",
     "I can plan tenures in my preferred unit and see correct tick marks",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "Medium",
     "US_007", "US_009",
     "1. Capture tenure scale signature\n2. Toggle to Mo, capture again — must differ\n3. Toggle back to Yr — signature returns to original"],

    ["US_009", "Reuse UI Validation Across Sub-Calculators", "FI_003",
     "QA tester",
     "Re-apply the same UI validation on Loan Amount and Loan Tenure calculators",
     "Quality bar is consistent across all three sub-calculators",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "Medium",
     "US_008", "US_010",
     "1. Run validateInputs() helper on Loan Amount Calculator — passes\n2. Run validateInputs() helper on Loan Tenure Calculator — passes"],

    ["US_010", "Parallel Multi-Browser Execution", "FI_004",
     "QA Lead",
     "Run the entire test suite in parallel on Chrome and Edge",
     "Both browsers are covered without doubling the wall-clock time",
     "Team SKRUM", "2026-05-19", "Haradhan Pal (Trainer)", "High",
     "US_009", "",
     "1. testng.xml parallel='tests'\n2. Chrome and Edge runners each run all 10 scenarios\n3. Firefox kept commented per requirement (point p)"],
]

ws1 = wb.create_sheet("UserStories")
ws1.sheet_view.showGridLines = False
# Title row
ncols = len(us_headers)
ws1.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
t = ws1.cell(row=1, column=1, value="User Story Document — EMI Calculator (INTQEA26QE003)")
t.font = title_font(14)
t.fill = PatternFill("solid", fgColor=BLUE_HEADER_FILL)
t.alignment = Alignment(horizontal="center", vertical="center")
ws1.row_dimensions[1].height = 30
# Headers (row 2)
for c, h in enumerate(us_headers, 1):
    ws1.cell(row=2, column=c, value=h)
style_blue_header_row(ws1, 2, ncols)
# Data (row 3+)
for r, row in enumerate(us_rows, start=3):
    for c, val in enumerate(row, start=1):
        ws1.cell(row=r, column=c, value=val)
    style_data_row(ws1, r, ncols)
set_widths(ws1, [12, 32, 18, 18, 35, 35, 18, 13, 22, 10, 18, 18, 60])
ws1.freeze_panes = "A3"


# =====================================================================
# 2) TestScenarios  (plain header style — matches reference)
# =====================================================================
ts_headers = ["Module", "Scenario ID", "Scenario Title", "Scenario Description", "Requirement ID"]
ts_rows = [
    ["Car_Loan", "EMI_TS001_CarEMIcalc",
     "Validate EMI for 15L car loan at 9.5% for 1 year",
     "Validate that the calculator on emicalculator.net computes the EMI correctly for the given inputs against the standard EMI formula.",
     "FR-001"],

    ["Car_Loan", "EMI_TS002_FirstMonthInterest",
     "Verify first month interest amount",
     "Validate that the first month's interest equals Principal × monthly rate, by reading the expanded year row in the schedule table.",
     "FR-002"],

    ["Car_Loan", "EMI_TS003_FirstMonthPrincipal",
     "Verify first month principal amount",
     "Validate that the first month's principal equals (EMI − first month interest).",
     "FR-002"],

    ["Car_Loan", "EMI_TS004_CarExcelExport",
     "Export car loan EMI summary to Excel",
     "Validate that the EMI summary (principal, rate, tenure, EMI, first-month split, total interest) is written to an .xlsx file via Apache POI.",
     "FR-003"],

    ["Home_Loan", "EMI_TS005_MenuNavigation",
     "Navigate to Home Loan Calculator via top menu",
     "Validate that hovering 'Loan Calculators & Widgets' and clicking 'Home Loan EMI Calculator' opens the dedicated home-loan-emi-calculator page.",
     "FR-004"],

    ["Home_Loan", "EMI_TS006_YearlyExtract",
     "Extract year-on-year schedule and store in Excel",
     "Validate that the full year-on-year amortisation table is extracted (>=10 rows, 6 columns) and persisted to .xlsx via Apache POI.",
     "FR-005"],

    ["Loan_Calculator", "EMI_TS007_UIInputs",
     "EMI Calculator UI inputs and sliders operable",
     "Validate that on the EMI Calculator sub-tab, Amount/Interest/Tenure text boxes are enabled and all three sliders are rendered.",
     "FR-006"],

    ["Loan_Calculator", "EMI_TS008_TenureToggle",
     "Tenure unit toggle changes slider scale",
     "Validate that flipping the tenure unit between Yr and Mo changes the slider scale signature and reverts on flip-back.",
     "FR-007"],

    ["Loan_Calculator", "EMI_TS009_AmountReuse",
     "Reuse UI validation on Loan Amount Calculator",
     "Validate that the same validateInputs() helper passes on the Loan Amount Calculator sub-tab.",
     "FR-008"],

    ["Loan_Calculator", "EMI_TS010_TenureReuse",
     "Reuse UI validation on Loan Tenure Calculator",
     "Validate that the same validateInputs() helper passes on the Loan Tenure Calculator sub-tab.",
     "FR-008"],
]

ws2 = wb.create_sheet("TestScenarios")
ws2.sheet_view.showGridLines = False
for c, h in enumerate(ts_headers, 1):
    ws2.cell(row=1, column=c, value=h)
style_plain_header_row(ws2, 1, len(ts_headers))
for r, row in enumerate(ts_rows, start=2):
    for c, val in enumerate(row, start=1):
        ws2.cell(row=r, column=c, value=val)
    style_data_row(ws2, r, len(ts_headers))
set_widths(ws2, [16, 28, 44, 70, 14])
ws2.freeze_panes = "A2"


# =====================================================================
# 3) TestCases  (multi-row per TC, merged cells on shared columns)
# =====================================================================
tc_headers = [
    "GROUP", "TEST Case ID", "TEST DESCRIPTION", "DESIGNER", "TEST TYPE",
    "TEST DATA", "COMPLEXITY",
    "STEP #", "STEP DESCRIPTION", "STEP EXPECTED RESULTS",
    "Actual Result", "Status", "Defect",
]

# Per test case: (group, id, desc, designer, type, data, complexity, [(step#, desc, expected, actual, status, defect), ...])
test_cases = [
    ("Car_Loan", "FT_TC001_CarEMICalc",
     "Validate EMI calculation for 15L / 9.5% / 1yr against the EMI formula",
     "Team SKRUM", "UI + Functional",
     "Principal=1500000, Rate=9.5, Tenure=1 Year", "Medium", [
        (1, "Open https://emicalculator.net/", "EMI Calculator homepage loads", "Homepage loaded successfully", "PASS", "NA"),
        (2, "Click the 'Car Loan' tab", "Car Loan tab becomes active", "Car Loan tab activated", "PASS", "NA"),
        (3, "Enter Loan Amount = 1500000", "Loan amount field shows '15,00,000'", "Value entered correctly", "PASS", "NA"),
        (4, "Enter Interest = 9.5", "Interest field shows 9.5", "Value entered correctly", "PASS", "NA"),
        (5, "Enter Tenure = 1, unit = Yr", "Tenure shows 1 with Yr toggle active", "Value entered correctly", "PASS", "NA"),
        (6, "Read EMI from result tile", "EMI = Rs.1,31,524 (±Rs.2)", "EMI = Rs.1,31,525 — within tolerance", "PASS", "NA"),
     ]),
    ("Car_Loan", "FT_TC002_FirstMonthInterest",
     "Verify first month interest equals Principal x monthly rate",
     "Team SKRUM", "Functional",
     "Principal=1500000, Rate=9.5", "Medium", [
        (1, "Fill car loan inputs (15L/9.5%/1yr)", "EMI computed", "EMI displayed", "PASS", "NA"),
        (2, "Scroll to the schedule table", "Yearly schedule visible", "Schedule visible", "PASS", "NA"),
        (3, "Click year row for current year (2026) to expand", "Monthly table for 2026 visible", "Expanded", "PASS", "NA"),
        (4, "Read first month Interest cell", "Interest = Rs.11,875 (±Rs.2)", "Interest = Rs.11,875 — exact", "PASS", "NA"),
     ]),
    ("Car_Loan", "FT_TC003_FirstMonthPrincipal",
     "Verify first month principal = EMI − first month interest",
     "Team SKRUM", "Functional",
     "Principal=1500000, Rate=9.5, Tenure=1 Year", "Medium", [
        (1, "Fill car loan inputs (15L/9.5%/1yr)", "EMI computed", "EMI displayed", "PASS", "NA"),
        (2, "Expand the current year row in the schedule", "Monthly rows visible", "Expanded", "PASS", "NA"),
        (3, "Read first month Principal cell", "Principal = Rs.1,19,649 (±Rs.2)", "Principal = Rs.1,19,650 — within tolerance", "PASS", "NA"),
     ]),
    ("Car_Loan", "FT_TC004_CarExcelExport",
     "Export car loan summary to Excel and verify on disk",
     "Team SKRUM", "Integration",
     "Principal=1500000, Rate=9.5, Tenure=1 Year", "Medium", [
        (1, "Capture EMI, first month split, total interest", "Values read into context", "Captured", "PASS", "NA"),
        (2, "Call ExcelUtils.writeSheet(...) → output/CarLoan_EMI_Summary.xlsx", "Workbook created without IOException", "Workbook created", "PASS", "NA"),
        (3, "Read row count of 'CarLoanEMI' sheet", "Row count >= 2", "8 rows written", "PASS", "NA"),
     ]),
    ("Home_Loan", "FT_TC005_MenuNavigation",
     "Open Home Loan EMI Calculator from the top menu",
     "Team SKRUM", "UI Navigation",
     "Menu path: Loan Calculators & Widgets > Home Loan EMI Calculator", "Low", [
        (1, "Hover 'Loan Calculators & Widgets' menu", "Dropdown becomes visible", "Dropdown shown", "PASS", "NA"),
        (2, "Click 'Home Loan EMI Calculator' link", "Browser navigates to dedicated page", "Page loaded", "PASS", "NA"),
        (3, "Assert current URL contains 'home-loan-emi-calculator'", "URL match passes", "URL matched", "PASS", "NA"),
     ]),
    ("Home_Loan", "FT_TC006_YearlyExtract",
     "Extract year-on-year payment schedule and write to Excel",
     "Team SKRUM", "Functional + Integration",
     "Principal=2500000, Rate=8.5, Tenure=20 years", "High", [
        (1, "Enter loan inputs and trigger calculation", "Schedule table renders", "Rendered", "PASS", "NA"),
        (2, "Scroll to year-on-year schedule, wait for rows", "All yearly rows visible", "21 rows visible", "PASS", "NA"),
        (3, "Iterate tr.yearlypaymentdetails rows, collect cells", "2D grid built with header + 21 rows", "Grid built", "PASS", "NA"),
        (4, "Assert at least 10 yearly rows", "Row count assertion passes", "Passed (21 >= 10)", "PASS", "NA"),
        (5, "Write grid to output/HomeLoan_YearlySchedule.xlsx", "Workbook saved", "File saved (4592 bytes)", "PASS", "NA"),
        (6, "Assert file exists and size > 0", "Both assertions pass", "Both passed", "PASS", "NA"),
     ]),
    ("Loan_Calculator", "FT_TC007_UISanity",
     "EMI Calculator sub-tab: text boxes enabled and sliders displayed",
     "Team SKRUM", "UI",
     "URL: emicalculator.net/loan-calculator/", "Low", [
        (1, "Open Loan Calculator and click 'EMI Calculator' sub-tab", "EMI Calculator tab is active", "Tab active", "PASS", "NA"),
        (2, "Check Loan Amount text box enabled", "isEnabled() returns true", "True", "PASS", "NA"),
        (3, "Check Interest text box enabled", "isEnabled() returns true", "True", "PASS", "NA"),
        (4, "Check Loan Tenure text box enabled", "isEnabled() returns true", "True", "PASS", "NA"),
        (5, "Check all three sliders displayed", "All slider handles visible", "3 handles visible", "PASS", "NA"),
     ]),
    ("Loan_Calculator", "FT_TC008_TenureToggle",
     "Tenure unit toggle (Yr <-> Mo) changes the slider scale",
     "Team SKRUM", "UI",
     "Default tenure unit = Yr", "Medium", [
        (1, "Capture tenure scale signature (A)", "Signature A captured", "'0|5|10|15|20|25|30|'", "PASS", "NA"),
        (2, "Click tenure unit = Mo", "Slider repaints with month ticks", "Repainted", "PASS", "NA"),
        (3, "Capture signature B, assert B != A", "Signatures differ", "'0|60|120|180|240|300|360|' (differs)", "PASS", "NA"),
        (4, "Click tenure unit = Yr", "Slider repaints back", "Repainted", "PASS", "NA"),
        (5, "Capture C, assert C == A", "Signature returns to original", "Matched A", "PASS", "NA"),
     ]),
    ("Loan_Calculator", "FT_TC009_AmountCalcReuse",
     "Re-use UI validation on Loan Amount Calculator sub-tab",
     "Team SKRUM", "UI Reuse",
     "URL: emicalculator.net/loan-calculator/", "Low", [
        (1, "Click 'Loan Amount Calculator' sub-tab", "Tab activates", "Activated", "PASS", "NA"),
        (2, "Run validateInputs() helper", "Returns UIValidationResult", "Result returned", "PASS", "NA"),
        (3, "Assert amount/interest/tenure boxes enabled", "Assertion passes", "Passed", "PASS", "NA"),
        (4, "Assert all sliders displayed", "Assertion passes", "Passed", "PASS", "NA"),
     ]),
    ("Loan_Calculator", "FT_TC010_TenureCalcReuse",
     "Re-use UI validation on Loan Tenure Calculator sub-tab",
     "Team SKRUM", "UI Reuse",
     "URL: emicalculator.net/loan-calculator/", "Low", [
        (1, "Click 'Loan Tenure Calculator' sub-tab", "Tab activates", "Activated", "PASS", "NA"),
        (2, "Run validateInputs() helper", "Returns UIValidationResult", "Result returned", "PASS", "NA"),
        (3, "Assert amount/interest/EMI boxes enabled", "Assertion passes", "Passed", "PASS", "NA"),
        (4, "Assert all sliders displayed", "Assertion passes", "Passed", "PASS", "NA"),
     ]),
]

ws3 = wb.create_sheet("TestCases")
ws3.sheet_view.showGridLines = False
for c, h in enumerate(tc_headers, 1):
    ws3.cell(row=1, column=c, value=h)
style_plain_header_row(ws3, 1, len(tc_headers))

row = 2
for (group, tcid, desc, designer, ttype, data, complexity, steps) in test_cases:
    start = row
    for (sno, sdesc, sexp, actual, status, defect) in steps:
        # Only fill shared columns on the first step row; merge later.
        if row == start:
            ws3.cell(row=row, column=1, value=group)
            ws3.cell(row=row, column=2, value=tcid)
            ws3.cell(row=row, column=3, value=desc)
            ws3.cell(row=row, column=4, value=designer)
            ws3.cell(row=row, column=5, value=ttype)
            ws3.cell(row=row, column=6, value=data)
            ws3.cell(row=row, column=7, value=complexity)
        ws3.cell(row=row, column=8,  value=sno)
        ws3.cell(row=row, column=9,  value=sdesc)
        ws3.cell(row=row, column=10, value=sexp)
        ws3.cell(row=row, column=11, value=actual)
        ws3.cell(row=row, column=12, value=status)
        ws3.cell(row=row, column=13, value=defect)
        style_data_row(ws3, row, len(tc_headers))
        row += 1
    # Merge the shared columns (1..7) across this test case's step rows
    if row - start > 1:
        for col in range(1, 8):
            ws3.merge_cells(start_row=start, start_column=col, end_row=row - 1, end_column=col)
            ws3.cell(row=start, column=col).alignment = Alignment(
                horizontal="left", vertical="top", wrap_text=True)

set_widths(ws3, [16, 30, 38, 16, 18, 28, 12, 8, 38, 36, 30, 10, 10])
ws3.freeze_panes = "A2"


# =====================================================================
# 4) RTM
# =====================================================================
rtm_headers = ["Serial no", "Requirement id", "Requirment description",
               "Test scenario id", "Test case id", "Defect id"]
rtm_rows = [
    [1,  "FR-001", "Car loan EMI shall match the EMI formula within Rs.2 tolerance",          "EMI_TS001_CarEMIcalc",          "FT_TC001_CarEMICalc",           ""],
    [2,  "FR-002", "First month interest shall equal Principal x monthly rate",               "EMI_TS002_FirstMonthInterest",  "FT_TC002_FirstMonthInterest",   ""],
    [3,  "FR-002", "First month principal shall equal EMI minus first month interest",         "EMI_TS003_FirstMonthPrincipal", "FT_TC003_FirstMonthPrincipal",  ""],
    [4,  "FR-003", "Car loan EMI summary shall be exportable to Excel via Apache POI",         "EMI_TS004_CarExcelExport",      "FT_TC004_CarExcelExport",       ""],
    [5,  "FR-004", "Home Loan Calculator shall be reachable from the top menu",                "EMI_TS005_MenuNavigation",      "FT_TC005_MenuNavigation",       ""],
    [6,  "FR-005", "Year-on-year payment schedule shall be extractable and storable to Excel", "EMI_TS006_YearlyExtract",       "FT_TC006_YearlyExtract",        "DEF_001"],
    [7,  "FR-006", "EMI Calculator sub-tab text boxes and sliders shall be operable",          "EMI_TS007_UIInputs",            "FT_TC007_UISanity",             ""],
    [8,  "FR-007", "Switching tenure Year/Month shall change the slider scale",                "EMI_TS008_TenureToggle",        "FT_TC008_TenureToggle",         ""],
    [9,  "FR-008", "Same UI validation shall be reusable across all 3 sub-calculators",        "EMI_TS009_AmountReuse",         "FT_TC009_AmountCalcReuse",      ""],
    [10, "FR-008", "Same UI validation shall be reusable across all 3 sub-calculators",        "EMI_TS010_TenureReuse",         "FT_TC010_TenureCalcReuse",      ""],
    [11, "FR-009", "Suite shall execute in parallel on Chrome and Edge (Firefox excluded)",    "Cross-cutting",                 "All TCs (testng.xml)",          ""],
    [12, "FR-010", "Failures shall capture a screenshot attached to the report",               "Cross-cutting",                 "Hooks.@After",                  ""],
]

ws4 = wb.create_sheet("RTM")
ws4.sheet_view.showGridLines = False
for c, h in enumerate(rtm_headers, 1):
    ws4.cell(row=1, column=c, value=h)
style_plain_header_row(ws4, 1, len(rtm_headers))
for r, row in enumerate(rtm_rows, start=2):
    for c, val in enumerate(row, start=1):
        ws4.cell(row=r, column=c, value=val)
    style_data_row(ws4, r, len(rtm_headers))
set_widths(ws4, [10, 14, 60, 32, 32, 12])
ws4.freeze_panes = "A2"


# =====================================================================
# 5) Defects
# =====================================================================
def_headers = ["Serial no.", "Defect id", "Description", "Reproducible (yes/no)",
               "Steps to reproduce", "Severity", "Priority",
               "Reported by", "Reported date", "Status", "Remarks"]
def_rows = [
    [1, "DEF_001",
     "Home Loan year-on-year schedule extraction occasionally times out on first run because ad blocks reflow the table position",
     "Intermittent",
     "1. Open https://emicalculator.net/home-loan-emi-calculator/ on slow network\n2. Fill 25L / 8.5 / 20 yr\n3. Immediately scroll to schedule and try to extract rows",
     "Medium", "Medium", "Team SKRUM", "13-05-2026", "Open",
     "Workaround: explicit wait + scrollBy(0,800) added in HomeLoanPage.extractYearlySchedule()"],

    [2, "DEF_002",
     "Tenure scale signature returns a different ordering on Edge vs Chrome under high zoom (>=150%)",
     "Yes",
     "1. Set browser zoom to 150%\n2. Open Loan Calculator > EMI tab\n3. Capture scale signature on each browser",
     "Low", "Low", "Team SKRUM", "14-05-2026", "Open",
     "Cosmetic — does not affect functional flow; deferred to next sprint"],

    [3, "DEF_003",
     "When tenure is entered as 'Mo' before selecting Car Loan tab, loan amount slider re-snaps to default",
     "Yes",
     "1. Open homepage (Home Loan default)\n2. Switch tenure to Mo, enter 24\n3. Switch to Car Loan tab\n4. Loan amount has reset to 3,00,000",
     "Low", "Low", "Team SKRUM", "15-05-2026", "Closed - Won't Fix",
     "Third-party site behaviour — out of scope"],

    [4, "DEF_004",
     "First-month row occasionally reads stale value if the schedule has not finished animating",
     "Intermittent",
     "1. Run TC02 in headed mode at fast speed\n2. Observe occasional 1-rupee mismatch in the first month interest",
     "Low", "Medium", "Team SKRUM", "16-05-2026", "Open",
     "Added explicit wait on visibility of the inner monthly tbody before reading"],
]

ws5 = wb.create_sheet("Defects")
ws5.sheet_view.showGridLines = False
for c, h in enumerate(def_headers, 1):
    ws5.cell(row=1, column=c, value=h)
style_plain_header_row(ws5, 1, len(def_headers))
for r, row in enumerate(def_rows, start=2):
    for c, val in enumerate(row, start=1):
        ws5.cell(row=r, column=c, value=val)
    style_data_row(ws5, r, len(def_headers))
set_widths(ws5, [10, 12, 50, 18, 50, 12, 12, 18, 14, 18, 40])
ws5.freeze_panes = "A2"

wb.save(OUT)
print(f"Saved {OUT}")
