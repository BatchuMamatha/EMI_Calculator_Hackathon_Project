"""Generate src/test/resources/testdata/TestData.xlsx containing all per-TC
inputs so the feature files / config / step-defs don't duplicate them."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT  = ROOT / "src" / "test" / "resources" / "testdata" / "TestData.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

HEADER_FILL = "FF2F78C4"
WHITE       = "FFFFFFFF"
thin = Side(border_style="thin", color="FFBFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

def style_header(ws, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[1].height = 28

def style_data(ws, nrows, ncols):
    for r in range(2, nrows + 2):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Aptos Narrow", size=11)
            cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
            cell.border = BORDER

def set_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def write_sheet(wb, name, headers, rows, widths):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)
    style_header(ws, len(headers))
    for r, row in enumerate(rows, start=2):
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    style_data(ws, len(rows), len(headers))
    set_widths(ws, widths)
    ws.freeze_panes = "A2"

wb = Workbook()
wb.remove(wb.active)

# CarLoan sheet
write_sheet(wb, "CarLoan",
    headers=["TestCaseID", "LoanAmount", "InterestRate", "Tenure", "TenureUnit", "VerificationType", "Year"],
    rows=[
        ["TC01", 1500000, 9.5, 1, "Yr", "emi",                   2026],
        ["TC02", 1500000, 9.5, 1, "Yr", "first_month_interest",  2026],
        ["TC03", 1500000, 9.5, 1, "Yr", "first_month_principal", 2026],
        ["TC04", 1500000, 9.5, 1, "Yr", "excel_export",          2026],
    ],
    widths=[12, 14, 14, 10, 12, 24, 8])

# HomeLoan sheet
write_sheet(wb, "HomeLoan",
    headers=["TestCaseID", "LoanAmount", "InterestRate", "Tenure", "MinRows", "MenuItem"],
    rows=[
        ["TC05", "", "", "", "",   "Home Loan EMI Calculator"],
        ["TC06", 2500000, 8.5, 20, 10, ""],
    ],
    widths=[12, 14, 14, 10, 10, 32])

# LoanCalculator sheet
write_sheet(wb, "LoanCalculator",
    headers=["TestCaseID", "SubTab"],
    rows=[
        ["TC07", "EMI Calculator"],
        ["TC08", "EMI Calculator"],
        ["TC09", "Loan Amount Calculator"],
        ["TC10", "Loan Tenure Calculator"],
    ],
    widths=[12, 30])

wb.save(OUT)
print(f"Saved {OUT}")
